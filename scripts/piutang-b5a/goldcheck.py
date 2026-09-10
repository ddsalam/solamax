#!/usr/bin/env python3
"""Read-only, preregistered B5a gold-check for customer balances.

The ledger is an append-only, hash-chained JSONL file. Production predictions
are sealed before any EasyMax workbook is parsed. The only external database
statements this tool permits are SELECTs inside an explicit READ ONLY
transaction, plus transaction-local set_config for the existing RLS policy.
"""

from __future__ import annotations

import argparse
import fcntl
import hashlib
import json
import os
import re
import subprocess
import sys
from contextlib import contextmanager
from datetime import date as Date
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Iterator
from urllib.parse import parse_qsl, unquote, urlsplit

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only in an incomplete operator env
    load_workbook = None


LEDGER_SCHEMA = "solamax-piutang-b5a-ledger/v1"
REQUIRED_REPORT_TITLE = "DAFTAR SALDO HUTANG PIUTANG"
FORBIDDEN_REPORT_TITLE = "LAPORAN PENJUALAN HARIAN"
SENTINEL = "B5A_JSON:"
SECTIONS = ("piutang_lokal", "piutang_online", "hutang_lokal")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
REPO_ROOT = Path(__file__).resolve().parents[2]
PRIVATE_LEDGER_ROOT = (REPO_ROOT / "verification-queries-results" / "piutang-b5a").resolve()
PSQL_TIMEOUT_SECONDS = 140
WRITE_SQL_RE = re.compile(
    r"\b(INSERT|UPDATE|DELETE|MERGE|CREATE|ALTER|DROP|TRUNCATE|GRANT|REVOKE|CALL|COPY)\b",
    re.IGNORECASE,
)


class GoldCheckError(RuntimeError):
    """A hard evidence or safety-gate failure."""


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def sha256_json(value: Any) -> str:
    return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds").replace("+00:00", "Z")


def parse_utc(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def load_config(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("schema") != "solamax-piutang-b5a-cases/v1":
        raise GoldCheckError("schema konfigurasi B5a tidak dikenal")
    if data.get("instance") != "solamax:asia-southeast2:solamax-pg":
        raise GoldCheckError("target wajib persis solamax:asia-southeast2:solamax-pg")
    if data.get("required_role") != "dashboard_ro":
        raise GoldCheckError("B5a wajib memakai role baca-saja dashboard_ro")
    system_identifier = str(data.get("expected_system_identifier", ""))
    if not system_identifier.isdigit():
        raise GoldCheckError("expected_system_identifier produksi wajib berupa angka")
    if data.get("required_domains") != ["masters", "piutang", "hutang"]:
        raise GoldCheckError("domain sumber B5a wajib masters, piutang, hutang")
    seen_units: set[int] = set()
    seen_cases: set[str] = set()
    for unit in data.get("cases", []):
        unit_id = unit.get("unit_id")
        unit_code = str(unit.get("unit_code", ""))
        if not isinstance(unit_id, int) or unit_id <= 0 or not unit_code.isdigit():
            raise GoldCheckError("unit_id/unit_code konfigurasi tidak sah")
        if unit_id in seen_units:
            raise GoldCheckError(f"unit_id duplikat: {unit_id}")
        seen_units.add(unit_id)
        dates = unit.get("dates", [])
        if len(dates) < 3 or len(set(dates)) != len(dates):
            raise GoldCheckError(f"{unit_code}: minimal tiga tanggal unik")
        for date in dates:
            if not DATE_RE.fullmatch(str(date)):
                raise GoldCheckError(f"tanggal tidak sah: {date}")
            datetime.strptime(date, "%Y-%m-%d")
            case_id = f"{unit_code}@{date}"
            if case_id in seen_cases:
                raise GoldCheckError(f"kasus duplikat: {case_id}")
            seen_cases.add(case_id)
        absent = unit.get("sections_expected_absent", [])
        if any(section not in SECTIONS for section in absent):
            raise GoldCheckError(f"seksi absent tidak dikenal untuk {unit_code}")
    if len(seen_units) < 2:
        raise GoldCheckError("B5a memerlukan minimal dua unit")
    if int(data.get("minimum_observation_gap_seconds", 0)) < 60:
        raise GoldCheckError("jarak observasi minimal tidak boleh kurang dari 60 detik")
    Decimal(str(data.get("comparison_tolerance_rupiah")))
    return data


def require_private_ledger(path: Path) -> None:
    if path.exists() and path.is_symlink():
        raise GoldCheckError("ledger tidak boleh berupa symlink")
    resolved = path.resolve()
    if PRIVATE_LEDGER_ROOT not in resolved.parents:
        raise GoldCheckError(
            "ledger produksi wajib berada di verification-queries-results/piutang-b5a/"
        )
    relative = resolved.relative_to(REPO_ROOT)
    ignored = subprocess.run(
        ["git", "check-ignore", "-q", "--", str(relative)],
        cwd=REPO_ROOT,
        check=False,
        timeout=10,
    )
    if ignored.returncode != 0:
        raise GoldCheckError("path ledger produksi tidak dilindungi .gitignore")


@contextmanager
def workflow_lock(path: Path) -> Iterator[None]:
    """Serialize every read-check-append transition for one evidence ledger."""
    require_private_ledger(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    lock_path = path.with_name(f"{path.name}.lock")
    if lock_path.exists() and lock_path.is_symlink():
        raise GoldCheckError("lock ledger tidak boleh berupa symlink")
    descriptor = os.open(lock_path, os.O_RDWR | os.O_CREAT | os.O_NOFOLLOW, 0o600)
    with os.fdopen(descriptor, "a+", encoding="utf-8") as handle:
        os.chmod(lock_path, 0o600)
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        yield


def verify_record(record: dict[str, Any], index: int, previous_hash: str | None) -> str:
    if record.get("schema") != LEDGER_SCHEMA:
        raise GoldCheckError(f"ledger record {index}: schema tidak dikenal")
    if record.get("sequence") != index:
        raise GoldCheckError(f"ledger record {index}: urutan rusak")
    if record.get("previous_hash") != previous_hash:
        raise GoldCheckError(f"ledger record {index}: rantai previous_hash rusak")
    supplied = record.get("record_hash")
    unsigned = {key: value for key, value in record.items() if key != "record_hash"}
    if supplied != sha256_json(unsigned):
        raise GoldCheckError(f"ledger record {index}: hash tidak cocok")
    return supplied


def verify_records(records: list[dict[str, Any]]) -> None:
    previous_hash: str | None = None
    for index, record in enumerate(records, start=1):
        previous_hash = verify_record(record, index, previous_hash)


def read_records(path: Path) -> list[dict[str, Any]]:
    try:
        handle = path.open(encoding="utf-8")
    except FileNotFoundError:
        return []
    records: list[dict[str, Any]] = []
    with handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_SH)
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                raise GoldCheckError(f"ledger baris {line_number}: baris kosong terlarang")
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise GoldCheckError(f"ledger baris {line_number}: JSON rusak") from exc
    verify_records(records)
    return records


def append_record(path: Path, kind: str, payload: dict[str, Any]) -> dict[str, Any]:
    require_private_ledger(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor = os.open(
        path,
        os.O_RDWR | os.O_CREAT | os.O_APPEND | os.O_NOFOLLOW,
        0o600,
    )
    with os.fdopen(descriptor, "a+", encoding="utf-8") as handle:
        os.chmod(path, 0o600)
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        handle.seek(0)
        previous_hash: str | None = None
        sequence = 0
        for sequence, line in enumerate(handle, start=1):
            if not line.strip():
                raise GoldCheckError(f"ledger baris {sequence}: baris kosong terlarang")
            try:
                existing = json.loads(line)
            except json.JSONDecodeError as exc:
                raise GoldCheckError(f"ledger baris {sequence}: JSON rusak") from exc
            previous_hash = verify_record(existing, sequence, previous_hash)
        record: dict[str, Any] = {
            "schema": LEDGER_SCHEMA,
            "sequence": sequence + 1,
            "kind": kind,
            "recorded_at_utc": utc_now(),
            "previous_hash": previous_hash,
            "payload": payload,
        }
        record["record_hash"] = sha256_json(record)
        handle.seek(0, os.SEEK_END)
        handle.write(canonical_json(record) + "\n")
        handle.flush()
        os.fsync(handle.fileno())
        return record


def libpq_connection_env(raw: str, proxy_host: str | None, proxy_port: int) -> dict[str, str]:
    parsed = urlsplit(raw)
    if parsed.scheme not in {"postgres", "postgresql"} or not parsed.username:
        raise GoldCheckError("DATABASE_URL PostgreSQL tidak sah")
    query = dict(parse_qsl(parsed.query, keep_blank_values=True))
    host = proxy_host or query.get("host") or parsed.hostname
    if not host:
        raise GoldCheckError("host PostgreSQL tidak tersedia")
    result = {
        "PGHOST": host,
        "PGPORT": str(proxy_port if proxy_host else (parsed.port or 5432)),
        "PGUSER": unquote(parsed.username),
        "PGDATABASE": parsed.path.lstrip("/"),
        "PGCONNECT_TIMEOUT": "10",
    }
    if parsed.password is not None:
        result["PGPASSWORD"] = unquote(parsed.password)
    if "sslmode" in query:
        result["PGSSLMODE"] = query["sslmode"]
    return result


def assert_read_only_sql(sql: str) -> None:
    without_literals = re.sub(r"'(?:''|[^'])*'", "''", sql)
    match = WRITE_SQL_RE.search(without_literals)
    if match:
        raise GoldCheckError(f"SQL mutasi ditolak: {match.group(1).upper()}")
    required = ("BEGIN TRANSACTION READ ONLY", "ROLLBACK", "set_config('app.unit_ids'")
    if any(fragment not in sql for fragment in required):
        raise GoldCheckError("SQL wajib memakai transaksi READ ONLY, RLS scope, dan ROLLBACK")


def psql_json(expression_sql: str, unit_ids: Iterable[int]) -> dict[str, Any]:
    raw_url = os.environ.get("DATABASE_URL")
    if not raw_url:
        raise GoldCheckError("DATABASE_URL belum tersedia")
    ids = sorted(set(int(value) for value in unit_ids))
    if not ids or any(value <= 0 for value in ids):
        raise GoldCheckError("RLS unit scope kosong/tidak sah")
    scope = ",".join(str(value) for value in ids)
    script = f"""
BEGIN TRANSACTION READ ONLY;
SELECT set_config('app.unit_ids', '{scope}', true);
SELECT '{SENTINEL}' || ({expression_sql})::text;
ROLLBACK;
"""
    assert_read_only_sql(script)
    proxy_host = os.environ.get("B5A_PROXY_HOST") or None
    proxy_port = int(os.environ.get("B5A_PROXY_PORT", "5432"))
    env = os.environ.copy()
    env.pop("DATABASE_URL", None)
    env.update(libpq_connection_env(raw_url, proxy_host, proxy_port))
    env["PGOPTIONS"] = (
        "-c default_transaction_read_only=on -c statement_timeout=120000 "
        "-c lock_timeout=5000 -c application_name=solamax_piutang_b5a_readonly"
    )
    command = [
        os.environ.get("PSQL", "psql"),
        "-X",
        "--no-psqlrc",
        "-qAt",
        "-v",
        "ON_ERROR_STOP=1",
    ]
    completed = subprocess.run(
        command,
        input=script,
        text=True,
        capture_output=True,
        env=env,
        check=False,
        timeout=PSQL_TIMEOUT_SECONDS,
    )
    if completed.returncode != 0:
        detail = completed.stderr.strip().splitlines()[-1] if completed.stderr.strip() else "tanpa detail"
        raise GoldCheckError(f"query read-only gagal: {detail}")
    payloads = [line[len(SENTINEL):] for line in completed.stdout.splitlines() if line.startswith(SENTINEL)]
    if len(payloads) != 1:
        raise GoldCheckError("psql tidak memulangkan tepat satu payload B5A_JSON")
    return json.loads(payloads[0])


def guard_sql(required_role: str) -> str:
    role = required_role.replace("'", "''")
    return f"""
SELECT json_build_object(
  'system_identifier', (SELECT system_identifier::text FROM pg_control_system()),
  'database', current_database(),
  'role', current_user,
  'transaction_read_only', current_setting('transaction_read_only'),
  'rolsuper', r.rolsuper,
  'rolbypassrls', r.rolbypassrls,
  'public_schema_create', has_schema_privilege(current_user, 'public', 'CREATE'),
  'bppiut_insert', has_table_privilege(current_user, 'public.bppiut', 'INSERT'),
  'bppiut_update', has_table_privilege(current_user, 'public.bppiut', 'UPDATE'),
  'bppiut_delete', has_table_privilege(current_user, 'public.bppiut', 'DELETE'),
  'bphut_insert', has_table_privilege(current_user, 'public.bphut', 'INSERT'),
  'bphut_update', has_table_privilege(current_user, 'public.bphut', 'UPDATE'),
  'bphut_delete', has_table_privilege(current_user, 'public.bphut', 'DELETE'),
  'pelanggan_master_insert', has_table_privilege(current_user, 'public.pelanggan_master', 'INSERT'),
  'pelanggan_master_update', has_table_privilege(current_user, 'public.pelanggan_master', 'UPDATE'),
  'pelanggan_master_delete', has_table_privilege(current_user, 'public.pelanggan_master', 'DELETE'),
  'sync_state_insert', has_table_privilege(current_user, 'public.sync_state', 'INSERT'),
  'sync_state_update', has_table_privilege(current_user, 'public.sync_state', 'UPDATE'),
  'sync_state_delete', has_table_privilege(current_user, 'public.sync_state', 'DELETE'),
  'role_matches', current_user = '{role}'
) AS value FROM pg_roles r WHERE r.rolname = current_user
"""


def verify_guard(guard: dict[str, Any], config: dict[str, Any]) -> None:
    expected = {
        "system_identifier": config["expected_system_identifier"],
        "database": "solamax",
        "role": config["required_role"],
        "transaction_read_only": "on",
        "rolsuper": False,
        "rolbypassrls": False,
        "public_schema_create": False,
        "bppiut_insert": False,
        "bppiut_update": False,
        "bppiut_delete": False,
        "bphut_insert": False,
        "bphut_update": False,
        "bphut_delete": False,
        "pelanggan_master_insert": False,
        "pelanggan_master_update": False,
        "pelanggan_master_delete": False,
        "sync_state_insert": False,
        "sync_state_update": False,
        "sync_state_delete": False,
        "role_matches": True,
    }
    differences = {key: {"expected": value, "actual": guard.get(key)} for key, value in expected.items() if guard.get(key) != value}
    if differences:
        raise GoldCheckError(f"guard koneksi baca-saja gagal: {canonical_json(differences)}")


def config_unit_ids(config: dict[str, Any]) -> list[int]:
    return [int(unit["unit_id"]) for unit in config["cases"]]


def values_sql(config: dict[str, Any]) -> str:
    values: list[str] = []
    for unit in config["cases"]:
        name = str(unit["unit_name"]).replace("'", "''")
        code = str(unit["unit_code"])
        for date in unit["dates"]:
            values.append(f"({unit['unit_id']},'{code}','{name}',DATE '{date}')")
    return ",\n    ".join(values)


def sync_payload_sql(config: dict[str, Any], include_predictions: bool) -> str:
    ids = ",".join(str(value) for value in config_unit_ids(config))
    domains = ",".join("'" + value.replace("'", "''") + "'" for value in config["required_domains"])
    guard = guard_sql(config["required_role"])
    prediction_fields = ""
    prediction_ctes = ""
    if include_predictions:
        prediction_ctes = f""",
case_dates(unit_id, unit_code, unit_name, as_of_date) AS (
  VALUES
    {values_sql(config)}
),
piut_candidates AS (
  SELECT c.unit_id,
         c.unit_code,
         c.unit_name,
         c.as_of_date,
         trim(m.ckdplg) AS customer_code,
         nullif(trim(m.vcnmplg), '') AS customer_name,
         CASE WHEN position('.' in trim(m.ckdplg)) > 0
              THEN 'piutang_online' ELSE 'piutang_lokal' END AS section
    FROM case_dates c
    JOIN public.pelanggan_master m ON m.unit_id = c.unit_id
   WHERE position('.' in trim(m.ckdplg)) > 0
      OR (m.sjenis IN (1, 5) AND position('.' in trim(m.ckdplg)) = 0)
  UNION
  SELECT c.unit_id,
         c.unit_code,
         c.unit_name,
         c.as_of_date,
         trim(b.ckdplg) AS customer_code,
         max(nullif(trim(m.vcnmplg), '')) AS customer_name,
         'piutang_online'::text AS section
    FROM case_dates c
    JOIN public.bppiut b
      ON b.unit_id = c.unit_id
     AND b.dtgl <= c.as_of_date
     AND coalesce(b.sbatal, 0) = 0
     AND position('.' in trim(b.ckdplg)) > 0
    LEFT JOIN public.pelanggan_master m
      ON m.unit_id = b.unit_id
     AND trim(m.ckdplg) = trim(b.ckdplg)
   GROUP BY c.unit_id, c.unit_code, c.unit_name, c.as_of_date, trim(b.ckdplg)
),
piut_rows AS (
  SELECT p.unit_id,
         p.unit_code,
         p.unit_name,
         p.as_of_date,
         p.customer_code,
         max(p.customer_name) AS customer_name,
         p.section,
         coalesce(sum(CASE WHEN b.sjnsbp = 1 THEN coalesce(b.njumlah, 0) ELSE 0 END), 0) AS debit,
         coalesce(sum(CASE WHEN b.sjnsbp = 2 THEN coalesce(b.njumlah, 0) ELSE 0 END), 0) AS credit,
         count(b.ckdbppiut)::int AS ledger_rows
    FROM piut_candidates p
    LEFT JOIN public.bppiut b
      ON b.unit_id = p.unit_id
     AND trim(b.ckdplg) = p.customer_code
     AND b.dtgl <= p.as_of_date
     AND coalesce(b.sbatal, 0) = 0
   GROUP BY p.unit_id, p.unit_code, p.unit_name, p.as_of_date,
            p.customer_code, p.section
),
hut_rows AS (
  SELECT c.unit_id,
         c.unit_code,
         c.unit_name,
         c.as_of_date,
         trim(h.ckdplg) AS customer_code,
         max(nullif(trim(m.vcnmplg), '')) AS customer_name,
         'hutang_lokal'::text AS section,
         sum(CASE WHEN h.sjnsbp = 1 THEN coalesce(h.njumlah, 0) ELSE 0 END) AS debit,
         sum(CASE WHEN h.sjnsbp = 2 THEN coalesce(h.njumlah, 0) ELSE 0 END) AS credit,
         count(*)::int AS ledger_rows
    FROM case_dates c
    JOIN public.bphut h
      ON h.unit_id = c.unit_id
     AND h.dtgl <= c.as_of_date
     AND coalesce(h.sbatal, 0) = 0
    LEFT JOIN public.pelanggan_master m
      ON m.unit_id = h.unit_id
     AND trim(m.ckdplg) = trim(h.ckdplg)
   GROUP BY c.unit_id, c.unit_code, c.unit_name, c.as_of_date, trim(h.ckdplg)
),
predictions AS (
  SELECT unit_id, unit_code, unit_name, as_of_date, section, customer_code,
         customer_name, debit, credit, debit - credit AS saldo, ledger_rows
    FROM piut_rows
  UNION ALL
  SELECT unit_id, unit_code, unit_name, as_of_date, section, customer_code,
         customer_name, debit, credit, debit - credit AS saldo, ledger_rows
    FROM hut_rows
),
prediction_anomalies AS (
  SELECT 'bppiut_unsupported_sjnsbp'::text AS kind, count(*)::int AS rows
    FROM public.bppiut b
   WHERE b.unit_id IN ({ids}) AND coalesce(b.sbatal,0)=0
     AND (b.sjnsbp IS NULL OR b.sjnsbp NOT IN (1,2)) AND coalesce(b.njumlah,0) <> 0
  UNION ALL
  SELECT 'bphut_unsupported_sjnsbp', count(*)::int
    FROM public.bphut h
   WHERE h.unit_id IN ({ids}) AND coalesce(h.sbatal,0)=0
     AND (h.sjnsbp IS NULL OR h.sjnsbp NOT IN (1,2)) AND coalesce(h.njumlah,0) <> 0
  UNION ALL
  SELECT 'bphut_blank_customer', count(*)::int
    FROM public.bphut h
   WHERE h.unit_id IN ({ids}) AND coalesce(h.sbatal,0)=0
     AND coalesce(trim(h.ckdplg),'') = '' AND coalesce(h.njumlah,0) <> 0
)"""
        prediction_fields = """,
    'predictions', coalesce((
      SELECT json_agg(json_build_object(
        'case_id', unit_code || '@' || to_char(as_of_date,'YYYY-MM-DD'),
        'unit_id', unit_id,
        'unit_code', unit_code,
        'unit_name', unit_name,
        'date', to_char(as_of_date,'YYYY-MM-DD'),
        'section', section,
        'customer_code', customer_code,
        'customer_name', customer_name,
        'debit', debit::text,
        'credit', credit::text,
        'saldo', saldo::text,
        'ledger_rows', ledger_rows
      ) ORDER BY unit_id, as_of_date, section, customer_code)
      FROM predictions
    ), '[]'::json),
    'anomalies', coalesce((
      SELECT json_agg(json_build_object('kind', kind, 'rows', rows) ORDER BY kind)
      FROM prediction_anomalies
    ), '[]'::json)"""
    return f"""
WITH sync_rows AS (
  SELECT s.unit_id,
         s.domain,
         to_char(s.last_run_at AT TIME ZONE 'UTC','YYYY-MM-DD"T"HH24:MI:SS.US"Z"') AS last_run_at_utc,
         s.last_row_count,
         s.last_watermark
    FROM public.sync_state s
   WHERE s.unit_id IN ({ids}) AND s.domain IN ({domains})
),
connection_guard AS ({guard})
{prediction_ctes}
SELECT json_build_object(
  'observed_at_utc', to_char(clock_timestamp() AT TIME ZONE 'UTC','YYYY-MM-DD"T"HH24:MI:SS.US"Z"'),
  'guard', (SELECT value FROM connection_guard),
  'sync_rows', coalesce((
    SELECT json_agg(json_build_object(
      'unit_id', unit_id,
      'domain', domain,
      'last_run_at_utc', last_run_at_utc,
      'last_row_count', last_row_count,
      'last_watermark', last_watermark
    ) ORDER BY unit_id, domain) FROM sync_rows
  ), '[]'::json)
  {prediction_fields}
)
"""


def verify_sync_rows(payload: dict[str, Any], config: dict[str, Any]) -> None:
    expected = {
        (int(unit["unit_id"]), domain)
        for unit in config["cases"]
        for domain in config["required_domains"]
    }
    actual = {(int(row["unit_id"]), str(row["domain"])) for row in payload["sync_rows"]}
    if actual != expected:
        missing = sorted(expected - actual)
        extra = sorted(actual - expected)
        raise GoldCheckError(f"sync_state tidak lengkap; missing={missing}, extra={extra}")
    if any(row.get("last_run_at_utc") is None for row in payload["sync_rows"]):
        raise GoldCheckError("sync_state memuat last_run_at kosong")


def observation_fingerprint(payload: dict[str, Any]) -> str:
    keys = ("unit_id", "domain", "last_run_at_utc", "last_row_count", "last_watermark")
    normalized = [{key: row.get(key) for key in keys} for row in payload["sync_rows"]]
    return sha256_json(normalized)


def sync_observation_payload(
    config: dict[str, Any], payload: dict[str, Any], fingerprint: str
) -> dict[str, Any]:
    return {
        "instance": config["instance"],
        "config_sha256": sha256_json(config),
        "observed_at_utc": payload["observed_at_utc"],
        "fingerprint": fingerprint,
        "sync_rows": payload["sync_rows"],
    }


def observe_sync(config_path: Path, ledger_path: Path) -> None:
    with workflow_lock(ledger_path):
        _observe_sync_locked(config_path, ledger_path)


def _observe_sync_locked(config_path: Path, ledger_path: Path) -> None:
    config = load_config(config_path)
    payload = psql_json(sync_payload_sql(config, include_predictions=False), config_unit_ids(config))
    verify_guard(payload["guard"], config)
    verify_sync_rows(payload, config)
    record = append_record(
        ledger_path,
        "sync_observation",
        sync_observation_payload(config, payload, observation_fingerprint(payload)),
    )
    print(f"observasi sync append-only tersimpan: sequence={record['sequence']} fingerprint={record['payload']['fingerprint'][:12]}")


def case_map(config: dict[str, Any]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for unit in config["cases"]:
        for date in unit["dates"]:
            result[f"{unit['unit_code']}@{date}"] = unit
    return result


def select_zero_controls(predictions: list[dict[str, Any]], config: dict[str, Any]) -> list[dict[str, Any]]:
    controls: list[dict[str, Any]] = []
    by_unit: dict[str, list[dict[str, Any]]] = {}
    for row in predictions:
        by_unit.setdefault(str(row["unit_code"]), []).append(row)
    for unit in config["cases"]:
        code = str(unit["unit_code"])
        candidates: dict[tuple[str, str], list[dict[str, Any]]] = {}
        for row in by_unit.get(code, []):
            if Decimal(row["saldo"]) == 0:
                candidates.setdefault((row["section"], row["customer_code"]), []).append(row)
        if not candidates:
            continue
        ranked = sorted(
            candidates.values(),
            key=lambda rows: (
                -len({row["date"] for row in rows}),
                not any(Decimal(row["debit"]) != 0 or Decimal(row["credit"]) != 0 for row in rows),
                rows[0]["section"],
                rows[0]["customer_code"],
            ),
        )
        chosen = ranked[0]
        has_activity = any(Decimal(row["debit"]) != 0 or Decimal(row["credit"]) != 0 for row in chosen)
        controls.append(
            {
                "unit_code": code,
                "unit_name": unit["unit_name"],
                "section": chosen[0]["section"],
                "customer_code": chosen[0]["customer_code"],
                "dates": sorted({row["date"] for row in chosen}),
                "expected_saldo": "0",
                "reason": (
                    "ledger punya aktivitas, tetapi DEBET-KREDIT bersih nol"
                    if has_activity
                    else "pelanggan master sah tanpa mutasi ledger sampai tanggal pembanding"
                ),
            }
        )
    if not controls:
        raise GoldCheckError("tidak ada pelanggan bersaldo nol di seluruh set kasus")
    return controls


def verify_absent_controls(predictions: list[dict[str, Any]], config: dict[str, Any]) -> list[dict[str, Any]]:
    controls: list[dict[str, Any]] = []
    present = {(row["unit_code"], row["date"], row["section"]) for row in predictions}
    for unit in config["cases"]:
        for section in unit.get("sections_expected_absent", []):
            for date in unit["dates"]:
                key = (unit["unit_code"], date, section)
                if key in present:
                    raise GoldCheckError(f"kontrol ketiadaan gagal sebelum seal: {key}")
            controls.append(
                {
                    "unit_code": unit["unit_code"],
                    "unit_name": unit["unit_name"],
                    "section": section,
                    "dates": unit["dates"],
                    "expected": "section_not_printed",
                }
            )
    return controls


def seal_predictions(config_path: Path, ledger_path: Path) -> None:
    with workflow_lock(ledger_path):
        _seal_predictions_locked(config_path, ledger_path)


def _seal_predictions_locked(config_path: Path, ledger_path: Path) -> None:
    config = load_config(config_path)
    records = read_records(ledger_path)
    if any(record["kind"] == "prediction_seal" for record in records):
        raise GoldCheckError("prediction_seal sudah ada; append-only berarti tidak boleh disegel ulang")
    observations = [record for record in records if record["kind"] == "sync_observation"]
    if not observations:
        raise GoldCheckError("jalankan observe-sync lebih dulu")
    current = psql_json(sync_payload_sql(config, include_predictions=True), config_unit_ids(config))
    verify_guard(current["guard"], config)
    verify_sync_rows(current, config)
    previous = observations[-1]["payload"]
    current_config_hash = sha256_json(config)
    if previous.get("config_sha256") != current_config_hash:
        raise GoldCheckError("konfigurasi berubah sesudah observasi pertama")
    current_fingerprint = observation_fingerprint(current)
    gap = (parse_utc(current["observed_at_utc"]) - parse_utc(previous["observed_at_utc"])).total_seconds()
    minimum_gap = int(config["minimum_observation_gap_seconds"])
    if gap < minimum_gap:
        raise GoldCheckError(f"jarak observasi baru {gap:.1f} detik; minimum {minimum_gap} detik")
    if current_fingerprint != previous["fingerprint"]:
        append_record(
            ledger_path,
            "sync_observation",
            sync_observation_payload(config, current, current_fingerprint),
        )
        raise GoldCheckError("sync_state bergerak; observasi terbaru sudah di-append, tunggu lalu seal lagi")
    anomalies = [row for row in current["anomalies"] if int(row["rows"]) != 0]
    if anomalies:
        raise GoldCheckError(f"ledger memulangkan bentuk tak diprediksi: {canonical_json(anomalies)}")
    predictions = current["predictions"]
    known_cases = set(case_map(config))
    returned_cases = {row["case_id"] for row in predictions}
    unexpected_cases = sorted(returned_cases - known_cases)
    if unexpected_cases:
        raise GoldCheckError(f"sistem memulangkan kasus tak diprediksi: {unexpected_cases}")
    zero_controls = select_zero_controls(predictions, config)
    absent_controls = verify_absent_controls(predictions, config)
    second_observation = append_record(
        ledger_path,
        "sync_observation",
        sync_observation_payload(config, current, current_fingerprint),
    )
    seal = append_record(
        ledger_path,
        "prediction_seal",
        {
            "instance": config["instance"],
            "config_sha256": current_config_hash,
            "sync_observation_sequences": [observations[-1]["sequence"], second_observation["sequence"]],
            "sync_observation_gap_seconds": gap,
            "query_contract_sha256": sha256_json(sync_payload_sql(config, include_predictions=True)),
            "comparison_tolerance_rupiah": config["comparison_tolerance_rupiah"],
            "predictions": predictions,
            "zero_balance_controls": zero_controls,
            "absent_section_controls": absent_controls,
        },
    )
    print(f"prediksi tersegel append-only: sequence={seal['sequence']} hash={seal['record_hash']}")
    for control in zero_controls:
        print(f"kontrol nol {control['unit_name']}: {control['section']} {control['customer_code']}")
    for control in absent_controls:
        print(f"kontrol absent {control['unit_name']}: {control['section']} tidak dicetak")


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return re.sub(r"\s+", " ", text)


def compact_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalize_text(value))


def parse_decimal(value: Any, *, blank_zero: bool = True) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        if blank_zero:
            return Decimal(0)
        raise GoldCheckError("sel angka kosong")
    if isinstance(value, bool):
        raise GoldCheckError("boolean bukan angka saldo")
    if isinstance(value, (int, Decimal)):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = str(value).strip()
    if text.startswith("="):
        raise GoldCheckError("formula XLSX tidak memiliki cached value")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("Rp", "").replace("RP", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    elif text.count(".") > 1 or (text.count(".") == 1 and len(text.rsplit(".", 1)[1]) == 3):
        text = text.replace(".", "")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise GoldCheckError(f"angka tidak dapat dibaca: {value!r}") from exc
    return -result if negative else result


def detect_section(row: tuple[Any, ...]) -> str | None:
    joined = normalize_text(" ".join(str(value) for value in row if value is not None))
    if (
        "DAFTAR SALDO" not in joined
        or REQUIRED_REPORT_TITLE in joined
        or joined.startswith("TOTAL SALDO")
    ):
        return None
    if "PIUTANG" in joined and "ONLINE" in joined:
        return "piutang_online"
    if "PIUTANG" in joined and "LOKAL" in joined:
        return "piutang_lokal"
    if "HUTANG" in joined and REQUIRED_REPORT_TITLE not in joined:
        return "hutang_lokal"
    raise GoldCheckError(f"seksi DAFTAR SALDO tidak dikenal: {joined}")


def detect_columns(row: tuple[Any, ...]) -> dict[str, int] | None:
    result: dict[str, int] = {}
    for index, value in enumerate(row):
        header = compact_header(value)
        if header in {"KODE", "KD", "CKDPLG", "KODEPELANGGAN", "KDPELANGGAN"}:
            result.setdefault("code", index)
        elif header in {"NAMA", "NAMAPELANGGAN"}:
            result.setdefault("name", index)
        elif "DEBET" in header or "DEBIT" in header:
            result.setdefault("debit", index)
        elif "KREDIT" in header or "CREDIT" in header:
            result.setdefault("credit", index)
        elif "SALDO" in header:
            result.setdefault("saldo", index)
    return result if {"code", "debit", "credit", "saldo"}.issubset(result) else None


def date_tokens(value: str) -> set[str]:
    parsed = datetime.strptime(value, "%Y-%m-%d").date()
    return {
        parsed.strftime("%Y-%m-%d"),
        parsed.strftime("%d-%m-%Y"),
        parsed.strftime("%d/%m/%Y"),
        parsed.strftime("%d.%m.%Y"),
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_oracle(
    path: Path,
    *,
    expected_unit_code: str | None = None,
    expected_unit_name: str | None = None,
    expected_date: str | None = None,
) -> dict[str, Any]:
    if load_workbook is None:
        raise GoldCheckError("openpyxl belum terpasang")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise GoldCheckError("oracle wajib .xlsx/.xlsm")
    source_hash = sha256_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    required_title_seen = False
    forbidden_title_seen = False
    parsed: dict[str, dict[str, dict[str, str]]] = {}
    printed_sections: set[str] = set()
    section_totals: dict[str, dict[str, str]] = {}
    section_total_counts: dict[str, int] = {}
    online_marker_seen = False
    unit_marker_seen = expected_unit_code is None and expected_unit_name is None
    date_marker_seen = expected_date is None
    compact_unit_code = re.sub(r"\D", "", expected_unit_code or "")
    normalized_unit_name = normalize_text(expected_unit_name)
    expected_date_tokens = date_tokens(expected_date) if expected_date else set()
    metadata_open = True
    try:
        for sheet in workbook.worksheets:
            pending_section: str | None = None
            pending_age = 0
            active_section: str | None = None
            columns: dict[str, int] | None = None
            for row in sheet.iter_rows(values_only=True):
                row_text = normalize_text(" ".join(str(value) for value in row if value is not None))
                row_compact = re.sub(r"[^A-Z0-9]", "", row_text)
                required_title_seen = required_title_seen or REQUIRED_REPORT_TITLE in row_text
                forbidden_title_seen = forbidden_title_seen or FORBIDDEN_REPORT_TITLE in row_text
                online_marker_seen = online_marker_seen or "ONLINE" in row_text
                if metadata_open:
                    unit_marker_seen = unit_marker_seen or bool(
                        (compact_unit_code and compact_unit_code in row_compact)
                        or (normalized_unit_name and normalized_unit_name in row_text)
                    )
                    date_marker_seen = date_marker_seen or any(
                        isinstance(value, (Date, datetime))
                        and not isinstance(value, bool)
                        and value.strftime("%Y-%m-%d") == expected_date
                        for value in row
                        if value is not None
                    ) or any(token in row_text for token in expected_date_tokens)
                section = detect_section(row)
                if section:
                    metadata_open = False
                    pending_section = section
                    pending_age = 0
                    active_section = None
                    columns = None
                    continue
                if pending_section:
                    pending_age += 1
                    found = detect_columns(row)
                    if found:
                        active_section = pending_section
                        printed_sections.add(active_section)
                        parsed.setdefault(active_section, {})
                        columns = found
                        pending_section = None
                        continue
                    if pending_age > 12:
                        pending_section = None
                first_text = normalize_text(next((value for value in row if value is not None), ""))
                if not active_section or not columns:
                    if first_text.startswith("TOTAL SALDO"):
                        raise GoldCheckError("oracle memuat TOTAL SALDO tanpa seksi aktif")
                    continue
                if first_text.startswith("TOTAL SALDO"):
                    total = {
                        key: str(parse_decimal(row[index]))
                        for key, index in columns.items()
                        if key in {"debit", "credit", "saldo"}
                    }
                    section_total_counts[active_section] = section_total_counts.get(active_section, 0) + 1
                    if section_total_counts[active_section] != 1:
                        raise GoldCheckError(f"oracle {active_section}: TOTAL SALDO duplikat")
                    section_totals[active_section] = total
                    active_section = None
                    columns = None
                    continue
                code_index = columns["code"]
                code = normalize_text(row[code_index] if code_index < len(row) else None).replace(" ", "")
                if not code or code in {"KODE", "KD"}:
                    continue
                debit = parse_decimal(row[columns["debit"]] if columns["debit"] < len(row) else None)
                credit = parse_decimal(row[columns["credit"]] if columns["credit"] < len(row) else None)
                saldo = parse_decimal(row[columns["saldo"]] if columns["saldo"] < len(row) else None)
                if saldo != debit - credit:
                    raise GoldCheckError(
                        f"oracle {active_section}/{code}: SALDO {saldo} != DEBET-KREDIT {debit-credit}"
                    )
                section_rows = parsed[active_section]
                if code in section_rows:
                    raise GoldCheckError(f"oracle {active_section}: pelanggan duplikat {code}")
                name = ""
                if "name" in columns and columns["name"] < len(row):
                    name = str(row[columns["name"]] or "").strip()
                section_rows[code] = {
                    "customer_name": name,
                    "debit": str(debit),
                    "credit": str(credit),
                    "saldo": str(saldo),
                }
    finally:
        workbook.close()
    if forbidden_title_seen:
        raise GoldCheckError("oracle terlarang: Laporan Penjualan Harian")
    if not required_title_seen:
        raise GoldCheckError("judul oracle bukan DAFTAR SALDO HUTANG PIUTANG")
    if not unit_marker_seen:
        raise GoldCheckError("identitas unit oracle tidak cocok dengan case_id")
    if not date_marker_seen:
        raise GoldCheckError("tanggal oracle tidak cocok dengan case_id")
    for section, rows in parsed.items():
        debit = sum((Decimal(row["debit"]) for row in rows.values()), Decimal(0))
        credit = sum((Decimal(row["credit"]) for row in rows.values()), Decimal(0))
        saldo = sum((Decimal(row["saldo"]) for row in rows.values()), Decimal(0))
        if saldo != debit - credit:
            raise GoldCheckError(f"oracle {section}: ΣSALDO != Σ(DEBET-KREDIT)")
        printed_total = section_totals.get(section)
        if not printed_total:
            raise GoldCheckError(f"oracle {section}: TOTAL SALDO wajib tepat satu")
        if any(
            Decimal(printed_total[key]) != value
            for key, value in {"debit": debit, "credit": credit, "saldo": saldo}.items()
        ):
            raise GoldCheckError(f"oracle {section}: TOTAL SALDO cetak tidak cocok dengan Σ baris")
    if sha256_file(path) != source_hash:
        raise GoldCheckError("berkas oracle berubah saat sedang diparse")
    return {
        "printed_sections": sorted(printed_sections),
        "sections": parsed,
        "section_totals": section_totals,
        "online_marker_seen": online_marker_seen,
        "file_sha256": source_hash,
    }


def prediction_for_case(seal: dict[str, Any], case_id: str) -> dict[str, dict[str, dict[str, str]]]:
    result: dict[str, dict[str, dict[str, str]]] = {section: {} for section in SECTIONS}
    for row in seal["payload"]["predictions"]:
        if row["case_id"] != case_id:
            continue
        result[row["section"]][row["customer_code"]] = {
            "customer_name": row.get("customer_name") or "",
            "debit": row["debit"],
            "credit": row["credit"],
            "saldo": row["saldo"],
        }
    return result


def compare_case(
    expected: dict[str, dict[str, dict[str, str]]],
    oracle: dict[str, Any],
    absent_sections: set[str],
    tolerance: Decimal,
) -> dict[str, Any]:
    printed = set(oracle["printed_sections"])
    expected_printed = {section for section in SECTIONS if expected[section]}
    unexpected_sections = sorted(printed - expected_printed)
    if unexpected_sections:
        raise GoldCheckError(f"sistem memulangkan seksi tak diprediksi: {unexpected_sections}")
    if "piutang_online" in absent_sections and oracle.get("online_marker_seen"):
        raise GoldCheckError("oracle memuat penanda Online yang tidak berhasil diparse sebagai seksi")
    differences: list[dict[str, Any]] = []
    for section in SECTIONS:
        expected_rows = expected[section]
        actual_rows = oracle["sections"].get(section, {})
        if section in absent_sections:
            if expected_rows:
                raise GoldCheckError(f"prediksi internal kontradiktif: {section} ditandai absent tetapi punya baris")
            if section in printed:
                raise GoldCheckError(f"sistem memulangkan seksi tak diprediksi: {section}")
            continue
        if expected_rows and section not in printed:
            differences.append({"kind": "missing_section", "section": section})
            continue
        unexpected = sorted(set(actual_rows) - set(expected_rows))
        if unexpected:
            raise GoldCheckError(f"sistem memulangkan pelanggan tak diprediksi di {section}: {unexpected[:10]}")
        missing = sorted(set(expected_rows) - set(actual_rows))
        for customer_code in missing:
            differences.append({"kind": "missing_customer", "section": section, "customer_code": customer_code})
        for customer_code in sorted(set(expected_rows) & set(actual_rows)):
            for field in ("debit", "credit", "saldo"):
                predicted = Decimal(expected_rows[customer_code][field])
                actual = Decimal(actual_rows[customer_code][field])
                delta = actual - predicted
                if abs(delta) > tolerance:
                    differences.append(
                        {
                            "kind": "value_mismatch",
                            "section": section,
                            "customer_code": customer_code,
                            "field": field,
                            "predicted": str(predicted),
                            "oracle": str(actual),
                            "delta": str(delta),
                        }
                    )
    return {
        "status": "pass" if not differences else "fail",
        "differences": differences,
        "compared_customers": sum(len(rows) for rows in expected.values()),
        "compared_cells": 3 * sum(len(rows) for rows in expected.values()),
    }


def parse_oracle_argument(raw: str) -> tuple[str, Path]:
    if "=" not in raw:
        raise GoldCheckError("--oracle harus CASE_ID=/path/file.xlsx")
    case_id, path = raw.split("=", 1)
    if "@" not in case_id or not path:
        raise GoldCheckError("--oracle harus UNIT_CODE@YYYY-MM-DD=/path/file.xlsx")
    unit_code, date = case_id.split("@", 1)
    if not unit_code.isdigit() or not DATE_RE.fullmatch(date):
        raise GoldCheckError(f"case_id oracle tidak sah: {case_id}")
    return case_id, Path(path)


def compare_oracles(
    config_path: Path,
    ledger_path: Path,
    oracle_args: list[str],
    expected_seal_hash: str,
    acknowledged_failed: set[str],
) -> None:
    with workflow_lock(ledger_path):
        _compare_oracles_locked(
            config_path,
            ledger_path,
            oracle_args,
            expected_seal_hash,
            acknowledged_failed,
        )


def _compare_oracles_locked(
    config_path: Path,
    ledger_path: Path,
    oracle_args: list[str],
    expected_seal_hash: str,
    acknowledged_failed: set[str],
) -> None:
    config = load_config(config_path)
    records = read_records(ledger_path)
    seals = [record for record in records if record["kind"] == "prediction_seal"]
    if len(seals) != 1:
        raise GoldCheckError("wajib ada tepat satu prediction_seal sebelum oracle dibaca")
    seal = seals[0]
    if not re.fullmatch(r"[0-9a-f]{64}", expected_seal_hash.lower()):
        raise GoldCheckError("--expect-seal-hash wajib SHA-256 64 karakter")
    if seal["record_hash"] != expected_seal_hash.lower():
        raise GoldCheckError("seal ledger tidak cocok dengan hash preregistrasi eksternal")
    if seal["payload"]["config_sha256"] != sha256_json(config):
        raise GoldCheckError("config kasus tidak sama dengan config yang disegel")
    completed_cases = {
        record["payload"]["case_id"]
        for record in records
        if record["kind"] == "oracle_comparison" and record["payload"].get("status") == "pass"
    }
    failed_cases = {
        record["payload"]["case_id"]
        for record in records
        if record["kind"] == "oracle_comparison" and record["payload"].get("status") != "pass"
    }
    known_cases = case_map(config)
    invalid_acknowledgements = acknowledged_failed - failed_cases
    if invalid_acknowledgements:
        raise GoldCheckError(
            f"--acknowledge-failed tidak merujuk kegagalan tercatat: {sorted(invalid_acknowledgements)}"
        )
    oracle_hash_cases: dict[str, set[str]] = {}
    for record in records:
        if record["kind"] != "oracle_comparison":
            continue
        digest = record["payload"].get("oracle_sha256")
        if digest:
            oracle_hash_cases.setdefault(str(digest), set()).add(record["payload"]["case_id"])
    tolerance = Decimal(config["comparison_tolerance_rupiah"])
    for raw in oracle_args:
        case_id, path = parse_oracle_argument(raw)
        if case_id not in known_cases:
            raise GoldCheckError(f"sistem memulangkan kasus tak diprediksi: {case_id}")
        if case_id in completed_cases:
            raise GoldCheckError(f"hasil {case_id} sudah di-append; tidak boleh ditimpa")
        if case_id in failed_cases and case_id not in acknowledged_failed:
            raise GoldCheckError(
                f"{case_id} punya kegagalan terdahulu; STOP atau ulangi dengan --acknowledge-failed"
            )
        unit = known_cases[case_id]
        oracle = parse_oracle(
            path,
            expected_unit_code=str(unit["unit_code"]),
            expected_unit_name=str(unit["unit_name"]),
            expected_date=case_id.split("@", 1)[1],
        )
        reused_by = oracle_hash_cases.get(oracle["file_sha256"], set()) - {case_id}
        if reused_by:
            raise GoldCheckError(
                f"berkas oracle dipakai untuk case_id lain: {sorted(reused_by)}"
            )
        expected = prediction_for_case(seal, case_id)
        result = compare_case(
            expected,
            oracle,
            set(unit.get("sections_expected_absent", [])),
            tolerance,
        )
        append_record(
            ledger_path,
            "oracle_comparison",
            {
                "case_id": case_id,
                "unit_name": unit["unit_name"],
                "oracle_path_name": path.name,
                "oracle_sha256": oracle["file_sha256"],
                "oracle_printed_sections": oracle["printed_sections"],
                "prior_failure_acknowledged": case_id in acknowledged_failed,
                **result,
            },
        )
        oracle_hash_cases.setdefault(oracle["file_sha256"], set()).add(case_id)
        if result["status"] == "pass":
            completed_cases.add(case_id)
        print(f"{case_id}: {result['status']} ({result['compared_customers']} pelanggan, {result['compared_cells']} sel)")
        if result["status"] != "pass":
            raise GoldCheckError(f"{case_id}: {len(result['differences'])} perbedaan; berhenti")


def show_status(config_path: Path, ledger_path: Path) -> None:
    config = load_config(config_path)
    records = read_records(ledger_path)
    seals = [record for record in records if record["kind"] == "prediction_seal"]
    comparisons = [record for record in records if record["kind"] == "oracle_comparison"]
    expected_cases = set(case_map(config))
    done = {
        record["payload"]["case_id"]
        for record in comparisons
        if record["payload"].get("status") == "pass"
    }
    failed = [
        record["payload"]["case_id"]
        for record in comparisons
        if record["payload"].get("status") != "pass"
    ]
    print(f"ledger records: {len(records)}")
    print(f"prediction seal: {'ada' if len(seals) == 1 else 'belum ada'}")
    print(f"oracle compared: {len(done)}/{len(expected_cases)}")
    print(f"pending: {', '.join(sorted(expected_cases - done)) or '-'}")
    print(f"failed attempts: {len(failed)} across {len(set(failed))} case(s)")
    if failed:
        print(f"failed cases: {', '.join(sorted(set(failed)))}")
    if len(seals) == 1:
        print(f"prediction seal hash: {seals[0]['record_hash']}")
    if records:
        print(f"last hash: {records[-1]['record_hash']}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=Path(__file__).with_name("cases.json"))
    parser.add_argument("--ledger", type=Path, required=True)
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("observe-sync")
    subparsers.add_parser("seal-predictions")
    compare = subparsers.add_parser("compare")
    compare.add_argument("--oracle", action="append", required=True)
    compare.add_argument("--expect-seal-hash", required=True)
    compare.add_argument("--acknowledge-failed", action="append", default=[])
    subparsers.add_parser("status")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "observe-sync":
            observe_sync(args.config, args.ledger)
        elif args.command == "seal-predictions":
            seal_predictions(args.config, args.ledger)
        elif args.command == "compare":
            compare_oracles(
                args.config,
                args.ledger,
                args.oracle,
                args.expect_seal_hash,
                set(args.acknowledge_failed),
            )
        elif args.command == "status":
            show_status(args.config, args.ledger)
        else:  # pragma: no cover
            raise GoldCheckError(f"perintah tidak dikenal: {args.command}")
        return 0
    except (GoldCheckError, OSError, subprocess.SubprocessError, json.JSONDecodeError) as exc:
        print(f"STOP: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
