from __future__ import annotations

import importlib.util
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

from openpyxl import Workbook


MODULE_PATH = Path(__file__).with_name("goldcheck.py")
SPEC = importlib.util.spec_from_file_location("piutang_b5a_goldcheck", MODULE_PATH)
assert SPEC and SPEC.loader
goldcheck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(goldcheck)


def add_section(
    sheet,
    start_row: int,
    title: str,
    rows: list[tuple[str, str, float | int, float | int, float | int]],
    *,
    include_total: bool = True,
) -> int:
    sheet.cell(start_row, 1, title)
    headers = ["KODE PELANGGAN", "NAMA PELANGGAN", "DEBET", "KREDIT", "SALDO"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(start_row + 1, column, value)
    row_number = start_row + 2
    for values in rows:
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        row_number += 1
    debit = sum(row[2] for row in rows)
    credit = sum(row[3] for row in rows)
    saldo = sum(row[4] for row in rows)
    if include_total:
        sheet.cell(row_number, 1, f"TOTAL SALDO {title}")
        sheet.cell(row_number, 3, debit)
        sheet.cell(row_number, 4, credit)
        sheet.cell(row_number, 5, saldo)
        row_number += 1
    return row_number + 1


def write_oracle(
    path: Path,
    *,
    include_online: bool = False,
    wrong_identity: bool = False,
    include_local_total: bool = True,
    unknown_section: bool = False,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Saldo"
    sheet["A1"] = "DAFTAR SALDO HUTANG PIUTANG"
    sheet["A2"] = "SPBU 64.781.01 ADISUCIPTO - PER TANGGAL 01/09/2026"
    row_number = add_section(
        sheet,
        3,
        "DAFTAR SALDO PIUTANG PELANGGAN LOKAL",
        [
            ("PLG0001", "Aktif", 125, 25, 100),
            ("PLG0002", "Kontrol Nol", 50, 50, 1 if wrong_identity else 0),
        ],
        include_total=include_local_total,
    )
    if include_online:
        row_number = add_section(
            sheet,
            row_number,
            "DAFTAR SALDO PIUTANG PELANGGAN ONLINE",
            [("18.999.0010", "Online", 30, 10, 20)],
        )
    row_number = add_section(
        sheet,
        row_number,
        "DAFTAR SALDO HUTANG PELANGGAN LOKAL",
        [("PLG0099", "Supplier", 40, 70, -30)],
    )
    if unknown_section:
        sheet.cell(row_number, 1, "DAFTAR SALDO PIUTANG PELANGGAN MITRA")
        row_number += 2
    # The section phrases repeat in EasyMax's summary. They must not reset a
    # finished accumulator because no column header follows them.
    sheet.cell(row_number + 1, 1, "Summary")
    sheet.cell(row_number + 2, 1, "DAFTAR SALDO PIUTANG PELANGGAN LOKAL")
    sheet.cell(row_number + 2, 2, 100)
    sheet.cell(row_number + 3, 1, "DAFTAR SALDO HUTANG PELANGGAN LOKAL")
    sheet.cell(row_number + 3, 2, -30)
    workbook.save(path)


def write_config(path: Path) -> dict:
    config = {
        "schema": "solamax-piutang-b5a-cases/v1",
        "instance": "solamax:asia-southeast2:solamax-pg",
        "expected_system_identifier": "7650126488674766864",
        "required_role": "dashboard_ro",
        "required_domains": ["masters", "piutang", "hutang"],
        "minimum_observation_gap_seconds": 180,
        "comparison_tolerance_rupiah": "0.001",
        "cases": [
            {
                "unit_id": 3,
                "unit_code": "6478101",
                "unit_name": "Adisucipto",
                "dates": ["2026-09-01", "2026-09-04", "2026-09-09"],
                "sections_expected_absent": ["piutang_online"],
            },
            {
                "unit_id": 7,
                "unit_code": "63781002",
                "unit_name": "28 Oktober",
                "dates": ["2026-09-01", "2026-09-04", "2026-09-09"],
                "sections_expected_absent": [],
            },
        ],
    }
    path.write_text(json.dumps(config), encoding="utf-8")
    return config


def guard_payload(system_identifier: str = "7650126488674766864") -> dict:
    payload = {
        "system_identifier": system_identifier,
        "database": "solamax",
        "role": "dashboard_ro",
        "transaction_read_only": "on",
        "rolsuper": False,
        "rolbypassrls": False,
        "public_schema_create": False,
        "role_matches": True,
    }
    for table in ("bppiut", "bphut", "pelanggan_master", "sync_state"):
        for privilege in ("insert", "update", "delete"):
            payload[f"{table}_{privilege}"] = False
    return payload


def sync_payload(config: dict, observed_at: datetime, *, changed: bool = False) -> dict:
    sync_time = datetime(2026, 9, 10, 0, 0, tzinfo=timezone.utc)
    sync_rows = []
    for unit in config["cases"]:
        for domain in config["required_domains"]:
            sync_rows.append(
                {
                    "unit_id": unit["unit_id"],
                    "domain": domain,
                    "last_run_at_utc": (
                        sync_time + (timedelta(seconds=1) if changed and domain == "piutang" else timedelta())
                    ).isoformat().replace("+00:00", "Z"),
                    "last_row_count": 10,
                    "last_watermark": "stable",
                }
            )
    predictions = []
    for unit in config["cases"]:
        for value in unit["dates"]:
            predictions.append(
                {
                    "case_id": f"{unit['unit_code']}@{value}",
                    "unit_id": unit["unit_id"],
                    "unit_code": unit["unit_code"],
                    "unit_name": unit["unit_name"],
                    "date": value,
                    "section": "piutang_lokal",
                    "customer_code": f"PLG{unit['unit_id']:04d}",
                    "customer_name": "Kontrol",
                    "debit": "50",
                    "credit": "50",
                    "saldo": "0",
                }
            )
    return {
        "guard": guard_payload(),
        "observed_at_utc": observed_at.isoformat().replace("+00:00", "Z"),
        "sync_rows": sync_rows,
        "anomalies": [],
        "predictions": predictions,
    }


class OracleParserTests(unittest.TestCase):
    def test_parses_rows_totals_zero_and_absent_online_section(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path)
            parsed = goldcheck.parse_oracle(path)

        self.assertEqual(parsed["printed_sections"], ["hutang_lokal", "piutang_lokal"])
        self.assertNotIn("piutang_online", parsed["sections"])
        self.assertEqual(parsed["sections"]["piutang_lokal"]["PLG0002"]["saldo"], "0")
        self.assertEqual(parsed["section_totals"]["piutang_lokal"]["saldo"], "100")
        self.assertEqual(parsed["section_totals"]["hutang_lokal"]["saldo"], "-30")

    def test_parses_online_section_when_printed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path, include_online=True)
            parsed = goldcheck.parse_oracle(path)
        self.assertEqual(parsed["sections"]["piutang_online"]["18.999.0010"]["saldo"], "20")

    def test_rejects_wrong_per_customer_identity(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path, wrong_identity=True)
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "SALDO .* DEBET-KREDIT"):
                goldcheck.parse_oracle(path)

    def test_rejects_laporan_penjualan_harian(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrong.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "LAPORAN PENJUALAN HARIAN"
            workbook.save(path)
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "oracle terlarang"):
                goldcheck.parse_oracle(path)

    def test_binds_oracle_to_unit_and_date_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path)
            parsed = goldcheck.parse_oracle(
                path,
                expected_unit_code="6478101",
                expected_unit_name="Adisucipto",
                expected_date="2026-09-01",
            )
            self.assertEqual(parsed["sections"]["piutang_lokal"]["PLG0001"]["saldo"], "100")
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "tanggal oracle"):
                goldcheck.parse_oracle(
                    path,
                    expected_unit_code="6478101",
                    expected_unit_name="Adisucipto",
                    expected_date="2026-09-04",
                )

    def test_rejects_missing_section_total(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path, include_local_total=False)
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "TOTAL SALDO wajib"):
                goldcheck.parse_oracle(path)

    def test_rejects_unknown_saldo_section(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path, unknown_section=True)
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "seksi DAFTAR SALDO tidak dikenal"):
                goldcheck.parse_oracle(path)

    def test_rejects_duplicate_or_orphan_section_total(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "DAFTAR SALDO HUTANG PIUTANG"
            row_number = add_section(
                sheet,
                3,
                "DAFTAR SALDO PIUTANG PELANGGAN LOKAL",
                [("PLG0001", "Aktif", 125, 25, 100)],
            )
            sheet.cell(row_number, 1, "TOTAL SALDO DUPLIKAT")
            workbook.save(path)
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "TOTAL SALDO tanpa seksi aktif"):
                goldcheck.parse_oracle(path)


class ComparisonTests(unittest.TestCase):
    def setUp(self) -> None:
        self.expected = {
            "piutang_lokal": {
                "PLG0001": {"customer_name": "Aktif", "debit": "125", "credit": "25", "saldo": "100"},
                "PLG0002": {"customer_name": "Kontrol Nol", "debit": "50", "credit": "50", "saldo": "0"},
            },
            "piutang_online": {},
            "hutang_lokal": {
                "PLG0099": {"customer_name": "Supplier", "debit": "40", "credit": "70", "saldo": "-30"},
            },
        }

    def test_compares_every_customer_and_cell(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path)
            oracle = goldcheck.parse_oracle(path)
        result = goldcheck.compare_case(
            self.expected,
            oracle,
            {"piutang_online"},
            goldcheck.Decimal("0.001"),
        )
        self.assertEqual(result["status"], "pass")
        self.assertEqual(result["compared_customers"], 3)
        self.assertEqual(result["compared_cells"], 9)

    def test_unpredicted_customer_is_hard_stop(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path)
            oracle = goldcheck.parse_oracle(path)
        del self.expected["piutang_lokal"]["PLG0001"]
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "pelanggan tak diprediksi"):
            goldcheck.compare_case(
                self.expected,
                oracle,
                {"piutang_online"},
                goldcheck.Decimal("0.001"),
            )

    def test_printed_absent_section_is_hard_stop(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path, include_online=True)
            oracle = goldcheck.parse_oracle(path)
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "seksi tak diprediksi"):
            goldcheck.compare_case(
                self.expected,
                oracle,
                {"piutang_online"},
                goldcheck.Decimal("0.001"),
            )

    def test_unparsed_online_marker_is_hard_stop_for_absent_control(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path)
            oracle = goldcheck.parse_oracle(path)
        oracle["online_marker_seen"] = True
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "penanda Online"):
            goldcheck.compare_case(
                self.expected,
                oracle,
                {"piutang_online"},
                goldcheck.Decimal("0.001"),
            )

    def test_missing_customer_and_value_mismatch_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "saldo.xlsx"
            write_oracle(path)
            oracle = goldcheck.parse_oracle(path)
        expected = json.loads(json.dumps(self.expected))
        expected["piutang_lokal"]["PLG_MISSING"] = {
            "customer_name": "Missing",
            "debit": "1",
            "credit": "0",
            "saldo": "1",
        }
        expected["piutang_lokal"]["PLG0001"]["debit"] = "126"
        result = goldcheck.compare_case(
            expected,
            oracle,
            {"piutang_online"},
            goldcheck.Decimal("0.001"),
        )
        self.assertEqual(result["status"], "fail")
        self.assertTrue(any(item["kind"] == "missing_customer" for item in result["differences"]))
        self.assertTrue(
            any(
                item["kind"] == "value_mismatch" and item["field"] == "debit"
                for item in result["differences"]
            )
        )

    def test_comparison_tolerance_boundary(self) -> None:
        oracle = {
            "printed_sections": ["piutang_lokal", "hutang_lokal"],
            "sections": json.loads(json.dumps(self.expected)),
            "online_marker_seen": False,
        }
        oracle["sections"]["piutang_lokal"]["PLG0001"]["saldo"] = "100.001"
        at_boundary = goldcheck.compare_case(
            self.expected,
            oracle,
            {"piutang_online"},
            goldcheck.Decimal("0.001"),
        )
        self.assertEqual(at_boundary["status"], "pass")
        oracle["sections"]["piutang_lokal"]["PLG0001"]["saldo"] = "100.0011"
        over_boundary = goldcheck.compare_case(
            self.expected,
            oracle,
            {"piutang_online"},
            goldcheck.Decimal("0.001"),
        )
        self.assertEqual(over_boundary["status"], "fail")


class LedgerAndSafetyTests(unittest.TestCase):
    def test_append_only_hash_chain_detects_tampering(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ledger.jsonl"
            with mock.patch.object(goldcheck, "require_private_ledger"):
                first = goldcheck.append_record(path, "sync_observation", {"value": 1})
                second = goldcheck.append_record(path, "prediction_seal", {"value": 2})
            self.assertEqual(first["sequence"], 1)
            self.assertEqual(second["previous_hash"], first["record_hash"])
            records = goldcheck.read_records(path)
            records[0]["payload"]["value"] = 99
            path.write_text("\n".join(goldcheck.canonical_json(record) for record in records) + "\n")
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "hash tidak cocok"):
                goldcheck.read_records(path)

    def test_private_ledger_path_is_fixed_independent_of_cwd(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "verification-queries-results"):
                goldcheck.require_private_ledger(Path(directory) / "ledger.jsonl")
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "verification-queries-results"):
            goldcheck.require_private_ledger(goldcheck.REPO_ROOT / "scripts" / "ledger.jsonl")

    def test_workflow_lock_blocks_a_second_process(self) -> None:
        goldcheck.PRIVATE_LEDGER_ROOT.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=goldcheck.PRIVATE_LEDGER_ROOT) as directory:
            ledger = Path(directory) / "ledger.jsonl"
            script = (
                "import sys; from pathlib import Path; "
                "sys.path.insert(0, 'scripts/piutang-b5a'); import goldcheck; "
                "lock=goldcheck.workflow_lock(Path(sys.argv[1])); lock.__enter__(); "
                "print('acquired', flush=True); lock.__exit__(None, None, None)"
            )
            with goldcheck.workflow_lock(ledger):
                process = subprocess.Popen(
                    [sys.executable, "-c", script, str(ledger)],
                    cwd=goldcheck.REPO_ROOT,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                )
                with self.assertRaises(subprocess.TimeoutExpired):
                    process.communicate(timeout=0.2)
            stdout, stderr = process.communicate(timeout=5)
            self.assertEqual(process.returncode, 0, stderr)
            self.assertEqual(stdout.strip(), "acquired")

    def test_sql_guard_rejects_mutation(self) -> None:
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "SQL mutasi ditolak"):
            goldcheck.assert_read_only_sql(
                "BEGIN TRANSACTION READ ONLY; SELECT set_config('app.unit_ids','3',true); DELETE FROM x; ROLLBACK;"
            )

    def test_sql_guard_allows_privilege_names_inside_literals(self) -> None:
        goldcheck.assert_read_only_sql(
            "BEGIN TRANSACTION READ ONLY; SELECT set_config('app.unit_ids','3',true); "
            "SELECT has_table_privilege(current_user,'public.bppiut','INSERT'); ROLLBACK;"
        )

    def test_database_secret_is_split_into_libpq_environment(self) -> None:
        env = goldcheck.libpq_connection_env(
            "postgresql://dashboard_ro:s3cr%40t@localhost/solamax?schema=public&host=%2Fcloudsql%2Fprod",
            "127.0.0.1",
            5432,
        )
        self.assertEqual(env["PGHOST"], "127.0.0.1")
        self.assertEqual(env["PGUSER"], "dashboard_ro")
        self.assertEqual(env["PGPASSWORD"], "s3cr@t")
        self.assertEqual(env["PGDATABASE"], "solamax")

    def test_guard_rejects_staging_system_identifier_and_hutang_writes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "cases.json"
            config = write_config(config_path)
        staging = guard_payload("7659054651798528016")
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "system_identifier"):
            goldcheck.verify_guard(staging, config)
        writable_hutang = guard_payload()
        writable_hutang["bphut_update"] = True
        with self.assertRaisesRegex(goldcheck.GoldCheckError, "bphut_update"):
            goldcheck.verify_guard(writable_hutang, config)

    def test_psql_subprocess_is_bounded_and_secret_not_forwarded_as_url(self) -> None:
        result = subprocess.CompletedProcess(
            args=["psql"],
            returncode=0,
            stdout=f"{goldcheck.SENTINEL}{{\"ok\":true}}\n",
            stderr="",
        )
        with mock.patch.dict(
            os.environ,
            {"DATABASE_URL": "postgresql://dashboard_ro:secret@localhost/solamax"},
            clear=False,
        ), mock.patch.object(goldcheck.subprocess, "run", return_value=result) as run:
            self.assertEqual(goldcheck.psql_json("SELECT json_build_object('ok', true)", [3]), {"ok": True})
        self.assertEqual(run.call_args.kwargs["timeout"], goldcheck.PSQL_TIMEOUT_SECONDS)
        self.assertNotIn("DATABASE_URL", run.call_args.kwargs["env"])
        self.assertEqual(
            run.call_args.kwargs["env"]["PGOPTIONS"],
            "-c default_transaction_read_only=on -c statement_timeout=120000 "
            "-c lock_timeout=5000 -c application_name=solamax_piutang_b5a_readonly",
        )

    def test_zero_control_prefers_activity_and_records_available_dates(self) -> None:
        config = {
            "cases": [
                {
                    "unit_code": "6478101",
                    "unit_name": "Adisucipto",
                    "dates": ["2026-09-01", "2026-09-04", "2026-09-09"],
                }
            ]
        }
        predictions = [
            {
                "unit_code": "6478101",
                "date": date,
                "section": "piutang_lokal",
                "customer_code": "PLG0002",
                "debit": "50",
                "credit": "50",
                "saldo": "0",
            }
            for date in config["cases"][0]["dates"]
        ]
        controls = goldcheck.select_zero_controls(predictions, config)
        self.assertEqual(controls[0]["customer_code"], "PLG0002")
        self.assertEqual(controls[0]["dates"], config["cases"][0]["dates"])
        self.assertIn("aktivitas", controls[0]["reason"])


class SealWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.config_path = self.root / "cases.json"
        self.config = write_config(self.config_path)
        self.ledger_path = self.root / "ledger.jsonl"
        self.started = datetime(2026, 9, 10, 1, 0, tzinfo=timezone.utc)

    def append_initial_observation(self) -> dict:
        payload = sync_payload(self.config, self.started)
        fingerprint = goldcheck.observation_fingerprint(payload)
        goldcheck.append_record(
            self.ledger_path,
            "sync_observation",
            goldcheck.sync_observation_payload(self.config, payload, fingerprint),
        )
        return payload

    def test_seal_rejects_short_gap_without_appending(self) -> None:
        with mock.patch.object(goldcheck, "require_private_ledger"):
            self.append_initial_observation()
            current = sync_payload(self.config, self.started + timedelta(seconds=100))
            with mock.patch.object(goldcheck, "psql_json", return_value=current):
                with self.assertRaisesRegex(goldcheck.GoldCheckError, "jarak observasi"):
                    goldcheck.seal_predictions(self.config_path, self.ledger_path)
            self.assertEqual([row["kind"] for row in goldcheck.read_records(self.ledger_path)], ["sync_observation"])

    def test_seal_appends_moved_observation_and_stops(self) -> None:
        with mock.patch.object(goldcheck, "require_private_ledger"):
            self.append_initial_observation()
            current = sync_payload(self.config, self.started + timedelta(seconds=200), changed=True)
            with mock.patch.object(goldcheck, "psql_json", return_value=current):
                with self.assertRaisesRegex(goldcheck.GoldCheckError, "sync_state bergerak"):
                    goldcheck.seal_predictions(self.config_path, self.ledger_path)
            self.assertEqual(
                [row["kind"] for row in goldcheck.read_records(self.ledger_path)],
                ["sync_observation", "sync_observation"],
            )

    def test_stable_gap_creates_exactly_one_seal(self) -> None:
        with mock.patch.object(goldcheck, "require_private_ledger"):
            self.append_initial_observation()
            current = sync_payload(self.config, self.started + timedelta(seconds=200))
            with mock.patch.object(goldcheck, "psql_json", return_value=current):
                goldcheck.seal_predictions(self.config_path, self.ledger_path)
                with self.assertRaisesRegex(goldcheck.GoldCheckError, "prediction_seal sudah ada"):
                    goldcheck.seal_predictions(self.config_path, self.ledger_path)
            records = goldcheck.read_records(self.ledger_path)
            self.assertEqual(sum(row["kind"] == "prediction_seal" for row in records), 1)

    def test_seal_rejects_query_anomalies(self) -> None:
        with mock.patch.object(goldcheck, "require_private_ledger"):
            self.append_initial_observation()
            current = sync_payload(self.config, self.started + timedelta(seconds=200))
            current["anomalies"] = [{"kind": "unknown_section", "rows": 1}]
            with mock.patch.object(goldcheck, "psql_json", return_value=current):
                with self.assertRaisesRegex(goldcheck.GoldCheckError, "bentuk tak diprediksi"):
                    goldcheck.seal_predictions(self.config_path, self.ledger_path)
            self.assertFalse(
                any(row["kind"] == "prediction_seal" for row in goldcheck.read_records(self.ledger_path))
            )

    def test_seal_rejects_unexpected_case(self) -> None:
        with mock.patch.object(goldcheck, "require_private_ledger"):
            self.append_initial_observation()
            current = sync_payload(self.config, self.started + timedelta(seconds=200))
            unexpected = dict(current["predictions"][0])
            unexpected["case_id"] = "9999999@2026-09-01"
            unexpected["unit_code"] = "9999999"
            current["predictions"].append(unexpected)
            with mock.patch.object(goldcheck, "psql_json", return_value=current):
                with self.assertRaisesRegex(goldcheck.GoldCheckError, "kasus tak diprediksi"):
                    goldcheck.seal_predictions(self.config_path, self.ledger_path)
            self.assertFalse(
                any(row["kind"] == "prediction_seal" for row in goldcheck.read_records(self.ledger_path))
            )


class OracleWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.config_path = self.root / "cases.json"
        self.config = write_config(self.config_path)
        self.ledger_path = self.root / "ledger.jsonl"
        payload = sync_payload(
            self.config,
            datetime(2026, 9, 10, 1, 0, tzinfo=timezone.utc),
        )
        with mock.patch.object(goldcheck, "require_private_ledger"):
            self.seal = goldcheck.append_record(
                self.ledger_path,
                "prediction_seal",
                {
                    "config_sha256": goldcheck.sha256_json(self.config),
                    "predictions": payload["predictions"],
                },
            )

    @staticmethod
    def empty_oracle(digest: str = "a" * 64) -> dict:
        return {
            "printed_sections": ["piutang_lokal"],
            "sections": {"piutang_lokal": {}},
            "section_totals": {"piutang_lokal": {"debit": "0", "credit": "0", "saldo": "0"}},
            "online_marker_seen": False,
            "file_sha256": digest,
        }

    def test_failed_comparison_is_appended_then_stops(self) -> None:
        case_id = "6478101@2026-09-01"
        with mock.patch.object(goldcheck, "require_private_ledger"), mock.patch.object(
            goldcheck,
            "parse_oracle",
            return_value=self.empty_oracle(),
        ):
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "perbedaan; berhenti"):
                goldcheck.compare_oracles(
                    self.config_path,
                    self.ledger_path,
                    [f"{case_id}={self.root / 'oracle.xlsx'}"],
                    self.seal["record_hash"],
                    set(),
                )
        comparisons = [
            row for row in goldcheck.read_records(self.ledger_path) if row["kind"] == "oracle_comparison"
        ]
        self.assertEqual(len(comparisons), 1)
        self.assertEqual(comparisons[0]["payload"]["status"], "fail")
        self.assertTrue(any(item["kind"] == "missing_customer" for item in comparisons[0]["payload"]["differences"]))

    def test_compare_requires_external_seal_anchor(self) -> None:
        with mock.patch.object(goldcheck, "require_private_ledger"):
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "hash preregistrasi eksternal"):
                goldcheck.compare_oracles(
                    self.config_path,
                    self.ledger_path,
                    [f"6478101@2026-09-01={self.root / 'oracle.xlsx'}"],
                    "f" * 64,
                    set(),
                )

    def test_workbook_hash_cannot_be_reused_for_another_case(self) -> None:
        digest = "b" * 64
        with mock.patch.object(goldcheck, "require_private_ledger"):
            goldcheck.append_record(
                self.ledger_path,
                "oracle_comparison",
                {
                    "case_id": "6478101@2026-09-01",
                    "status": "fail",
                    "oracle_sha256": digest,
                },
            )
            with mock.patch.object(goldcheck, "parse_oracle", return_value=self.empty_oracle(digest)):
                with self.assertRaisesRegex(goldcheck.GoldCheckError, "case_id lain"):
                    goldcheck.compare_oracles(
                        self.config_path,
                        self.ledger_path,
                        [f"6478101@2026-09-04={self.root / 'same.xlsx'}"],
                        self.seal["record_hash"],
                        set(),
                    )

    def test_failed_case_requires_acknowledgement_and_status_shows_failure(self) -> None:
        case_id = "6478101@2026-09-01"
        with mock.patch.object(goldcheck, "require_private_ledger"):
            goldcheck.append_record(
                self.ledger_path,
                "oracle_comparison",
                {
                    "case_id": case_id,
                    "status": "fail",
                    "oracle_sha256": "c" * 64,
                },
            )
            with self.assertRaisesRegex(goldcheck.GoldCheckError, "--acknowledge-failed"):
                goldcheck.compare_oracles(
                    self.config_path,
                    self.ledger_path,
                    [f"{case_id}={self.root / 'retry.xlsx'}"],
                    self.seal["record_hash"],
                    set(),
                )
        output = io.StringIO()
        with redirect_stdout(output):
            goldcheck.show_status(self.config_path, self.ledger_path)
        self.assertIn("failed attempts: 1 across 1 case(s)", output.getvalue())
        self.assertIn(f"prediction seal hash: {self.seal['record_hash']}", output.getvalue())


if __name__ == "__main__":
    unittest.main()
